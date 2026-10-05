import html
import math
import zipfile

import ezdxf
import numpy as np
import pytest

from plumbline.core import crs as C
from plumbline.core.model import ImportBatch, SurveyPoint
from plumbline.core.project import Project
from plumbline.core.surface import build_surface_from_project, generate_contours
from plumbline.io import csv_points as CSV
from plumbline.io import dxf_io, gis_io, kml_io, landxml, reports

from test_core import make_site


# ============================================================ CSV
def test_csv_pnezd_with_spaces_in_description(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("1,6967000.5,2552000.25,501.10,EP B\n2,6967010.5,2552010.25,501.20,EP\n3,6967020,2552020,,TREE 18 OAK\n")
    sn = CSV.sniff(f)
    assert sn.delimiter == "," and not sn.has_header
    assert sn.roles[:5] == ["number", "northing", "easting", "elevation", "description"]
    b = CSV.read_points(f, CSV.CsvMapping(",", 0, sn.roles))
    assert [p.number for p in b.points] == ["1", "2", "3"]
    assert b.points[0].x == 2552000.25 and b.points[0].y == 6967000.5 and b.points[0].desc == "EP B"
    assert math.isnan(b.points[2].z) and b.points[2].desc == "TREE 18 OAK"


def test_csv_header_tab_and_headerless_space(tmp_path):
    f = tmp_path / "h.csv"
    f.write_text("Point\tX\tY\tElev\tDesc\n10\t1000.5\t2000.5\t12.5\tGS\n11\t1010.5\t2010.5\t12.7\tGS\n")
    sn = CSV.sniff(f)
    assert sn.delimiter == "\t" and sn.has_header
    assert sn.roles == ["number", "easting", "northing", "elevation", "description"]
    b = CSV.read_points(f, CSV.CsvMapping("\t", 1, sn.roles))
    assert (b.points[0].x, b.points[0].y) == (1000.5, 2000.5)
    g = tmp_path / "s.txt"
    g.write_text("1 6967000.5 2552000.25 501.1 TOP OF CURB\n2 6967010.5 2552010.25 501.2 EP\n")
    sn = CSV.sniff(g)
    assert sn.delimiter == " "
    b = CSV.read_points(g, CSV.CsvMapping(" ", 0, sn.roles))
    assert b.points[0].desc == "TOP OF CURB" and b.points[1].desc == "EP"


def test_csv_bad_rows_and_null_z_and_autonumber(tmp_path):
    f = tmp_path / "b.csv"
    f.write_text("100,200,-9999\nabc,def,1\n101,201,5\n")
    b = CSV.read_points(f, CSV.CsvMapping(",", 0, ["northing", "easting", "elevation"], null_z=-9999.0, z_scale=2.0))
    assert len(b.points) == 2 and math.isnan(b.points[0].z) and b.points[1].z == 10.0
    assert [p.number for p in b.points] == ["1", "2"] and any("skipped" in m for m in b.messages)
    with pytest.raises(ValueError):
        CSV.read_points(f, CSV.CsvMapping(",", 0, ["number", "elevation"]))


def test_csv_geographic_detect_and_ne_order(tmp_path):
    f = tmp_path / "g.csv"
    f.write_text("1,32.7668,-96.5992,150.0,GS\n2,32.7669,-96.5993,150.5,GS\n")
    sn = CSV.sniff(f)
    assert "latitude" in sn.roles and "longitude" in sn.roles
    b = CSV.read_points(f, CSV.CsvMapping(",", 0, sn.roles))
    assert b.info["geographic"] and b.points[0].x == -96.5992 and b.points[0].y == 32.7668
    # swapped N/E in a projected file is detected via the CRS area of use
    pr = C.ProjectCRS.from_epsg(2276)
    b = ImportBatch()
    b.points = [SurveyPoint(0, str(i), 6967000.0 + i, 2552000.0 + i, 1.0) for i in range(5)]      # x=northing (wrong)
    assert CSV.guess_ne_order(b, pr) == "swap"
    b.points = [SurveyPoint(0, str(i), 2552000.0 + i, 6967000.0 + i, 1.0) for i in range(5)]
    assert CSV.guess_ne_order(b, pr) == "ok"


def test_csv_write_roundtrip(tmp_path):
    pr = make_site(4)
    f = tmp_path / "o.csv"
    CSV.write_points_csv(f, pr.points.values())
    sn = CSV.sniff(f)
    assert sn.has_header
    b = CSV.read_points(f, CSV.CsvMapping(",", 1, sn.roles))
    assert len(b.points) == len(pr.points)
    p0, q0 = next(iter(pr.points.values())), b.points[0]
    assert (q0.x, q0.y, q0.z) == pytest.approx((p0.x, p0.y, p0.z), abs=5e-4)


# ============================================================ LandXML
SAMPLE_LANDXML = """<?xml version="1.0"?>
<LandXML xmlns="http://www.landxml.org/schema/LandXML-1.2" version="1.2">
 <Units><Imperial linearUnit="USSurveyFoot" areaUnit="squareFoot"/></Units>
 <CoordinateSystem name="NAD83 / Texas North Central (ftUS)" epsgCode="2276"/>
 <CgPoints><CgPoint name="1" code="EP" desc="EP B">6967000.5 2552000.25 501.1</CgPoint>
           <CgPoint name="2">6967010.5 2552010.25</CgPoint></CgPoints>
 <Surfaces><Surface name="EG"><Definition surfType="TIN">
   <Pnts><P id="1">0 0 10</P><P id="2">0 10 11</P><P id="3">10 10 12</P><P id="4">10 0 11</P></Pnts>
   <Faces><F>1 2 3</F><F>1 3 4</F><F i="1">1 2 4</F></Faces></Definition></Surface></Surfaces>
 <Parcels><Parcel name="Lot 1"><CoordGeom>
   <Line><Start>0 0</Start><End>0 100</End></Line>
   <Curve rot="cw" crvType="arc" radius="50"><Start>0 100</Start><Center>0 150</Center><End>50 150</End></Curve>
   <Line><Start>50 150</Start><End>0 0</End></Line></CoordGeom></Parcel></Parcels>
 <Alignments><Alignment name="CL" length="100"><CoordGeom><Line><Start>0 0</Start><End>100 0</End></Line></CoordGeom></Alignment></Alignments>
</LandXML>"""


def test_landxml_read(tmp_path):
    f = tmp_path / "s.xml"
    f.write_text(SAMPLE_LANDXML)
    b = landxml.read_landxml(f)
    assert b.info["units"] == "ftUS" and b.info["epsg"] == 2276
    assert len(b.points) == 2 and (b.points[0].x, b.points[0].y, b.points[0].z) == (2552000.25, 6967000.5, 501.1)
    assert b.points[0].desc == "EP B" and math.isnan(b.points[1].z)
    sf = b.surfaces[0]
    assert sf.name == "EG" and len(sf.pts) == 4 and len(sf.tris) == 2        # invisible face skipped
    assert sf.pts[1].tolist() == [10.0, 0.0, 11.0]                            # x=E, y=N
    par = [p for p in b.polylines if p.attrs.get("landxml") == "Parcel"][0]
    assert len(par.verts) == 3 and par.closed and par.bulges is not None
    assert any(p.attrs.get("landxml") == "Alignment" for p in b.polylines)


def test_landxml_roundtrip(tmp_path):
    pr = make_site(8)
    sf, _ = build_surface_from_project(pr, "EG", {})
    pr.add_surface(sf)
    pr.add_polyline([[0, 0, 0], [10, 0, 0], [10, 10, 0]], "BLDG", closed=True, bulges=[0, 0.5, 0], kind="parcel")
    f = tmp_path / "o.xml"
    st = landxml.write_landxml(f, pr)
    assert st["points"] == len(pr.points) and st["surfaces"] == 1 and st["polylines"] == 1
    b = landxml.read_landxml(f)
    assert len(b.points) == len(pr.points) and b.info["units"] == "ftUS" and b.info["epsg"] == 2276
    s2 = b.surfaces[0]
    assert len(s2.pts) == len(sf.pts) and len(s2.tris) == len(sf.tris)
    assert np.allclose(np.sort(s2.pts[:, 2]), np.sort(sf.pts[:, 2]), atol=1e-3)
    pl = b.polylines[0]
    assert pl.closed and pl.bulges is not None and pl.bulges[1] == pytest.approx(0.5, abs=1e-6)


def test_landxml_garbage_does_not_crash(tmp_path):
    f = tmp_path / "e.xml"
    f.write_text("<LandXML xmlns='http://www.landxml.org/schema/LandXML-1.2'></LandXML>")
    b = landxml.read_landxml(f)
    assert b.is_empty() and b.messages


# ============================================================ DXF
def test_dxf_roundtrip_geometry_layers_and_units(tmp_path):
    pr = make_site(5)
    pr.layers["POINTS"].color = (255, 0, 0)
    pr.add_polyline([[0, 0, 100], [10, 0, 100], [10, 10, 100]], "BLDG", closed=True, bulges=[0, 0.5, 0])
    pr.add_polyline([[0, 0, 100], [10, 0, 105]], "SLOPE")
    pr.add_text(5, 5, "NOTE", 2.5, 30.0)
    f = tmp_path / "o.dxf"
    st = dxf_io.write_dxf(f, pr, dxf_io.DxfOptions(text_height=2.0))
    assert st["points"] == 25 and st["polylines"] == 1 and st["polylines3d"] == 1
    doc = ezdxf.readfile(f)
    assert doc.header["$INSUNITS"] == 21
    assert "BLDG" in doc.layers and "PT-NUMBER" in doc.layers
    lw = [e for e in doc.modelspace().query("LWPOLYLINE")][0]
    assert lw.closed and lw.dxf.elevation == 100.0 and list(lw.get_points("xyb"))[1][2] == pytest.approx(0.5)
    assert doc.layers.get("POINTS").dxf.true_color == 0xFF0000
    # export must not alter the project
    assert "PT-NUMBER" not in pr.layers
    b = dxf_io.read_dxf(f)
    assert b.info["units"] == "ftUS" and len(b.points) == 25
    plines = {p.layer: p for p in b.polylines}
    assert plines["BLDG"].closed and plines["BLDG"].bulges[1] == pytest.approx(0.5)
    assert plines["SLOPE"].verts[1, 2] == pytest.approx(105.0)
    assert any(t.text == "NOTE" and t.rotation == pytest.approx(30.0) for t in b.texts)


def test_dxf_export_contours_faces_blocks_and_filters(tmp_path):
    pr = make_site(10)
    sf, _ = build_surface_from_project(pr, "S", {})
    pr.add_surface(sf)
    generate_contours(pr, sf, 1.0, 5, labels=True, text_height=2.0)
    f = tmp_path / "c.dxf"
    st = dxf_io.write_dxf(f, pr, dxf_io.DxfOptions(points_mode="block", include_surface_faces=True,
                                                   include_contour_labels=False))
    assert st["contours"] > 3 and st["faces"] == len(sf.tris)
    doc = ezdxf.readfile(f)
    ctr = [e for e in doc.modelspace().query("LWPOLYLINE") if e.dxf.layer.startswith("CONTOUR")]
    assert ctr and all(e.dxf.elevation == pytest.approx(round(e.dxf.elevation)) for e in ctr)
    ins = list(doc.modelspace().query("INSERT"))
    assert len(ins) == 100 and ins[0].get_attrib_text("NUMBER") == "1"
    assert not [t for t in doc.modelspace().query("TEXT") if t.dxf.layer == "CONTOUR-LABEL"]
    # DXF read-back recovers the TIN from 3DFACEs
    b = dxf_io.read_dxf(f)
    assert b.surfaces and len(b.surfaces[0].tris) == len(sf.tris)
    # hidden layers are skipped
    pr.layers["POINTS"].visible = False
    st = dxf_io.write_dxf(tmp_path / "d.dxf", pr)
    assert st["points"] == 0


def test_dxf_import_entity_zoo(tmp_path):
    doc = ezdxf.new("R2010")
    doc.header["$INSUNITS"] = 6
    msp = doc.modelspace()
    doc.layers.add("ROAD", color=3)
    msp.add_line((0, 0, 1), (10, 0, 2), dxfattribs={"layer": "ROAD"})
    msp.add_circle((50, 50), 5)
    msp.add_arc((0, 0), 10, 0, 90)
    msp.add_lwpolyline([(0, 0), (5, 5), (10, 0)])
    msp.add_polyline3d([(0, 0, 0), (1, 1, 1), (2, 0, 3)])
    msp.add_mtext("hello\\Pworld", dxfattribs={"insert": (1, 1), "char_height": 2})
    msp.add_spline([(0, 0), (5, 5), (10, 0)])
    msp.add_hatch()
    msp.add_point((1, 2, 3))
    blk = doc.blocks.new("TREE")
    blk.add_circle((0, 0), 1)
    blk.add_attdef("DESC", (0, 0))
    ins = msp.add_blockref("TREE", (7, 8, 9))
    ins.add_auto_attribs({"DESC": "OAK"})
    f = tmp_path / "z.dxf"
    doc.saveas(f)
    b = dxf_io.read_dxf(f)
    assert b.info["units"] == "m"
    kinds = len(b.polylines)
    assert kinds == 6                        # line, circle, arc, lwpoly, poly3d, spline
    assert any("HATCH" in m for m in b.messages)
    assert b.texts and "hello" in b.texts[0].text
    pts = {p.desc: p for p in b.points}
    assert pts["OAK"].z == 9.0 and pts[""].z == 3.0
    circ = [p for p in b.polylines if p.closed and p.bulges is not None][0]
    assert circ.bulges.tolist() == [1.0, 1.0]
    # exploding blocks turns the block circle into geometry instead of a point
    b2 = dxf_io.read_dxf(f, explode_blocks=True)
    assert len(b2.polylines) == kinds + 1


def test_dxf_all_zero_points_become_no_elevation(tmp_path):
    doc = ezdxf.new()
    doc.modelspace().add_point((1, 2, 0))
    doc.modelspace().add_point((3, 4, 0))
    doc.saveas(tmp_path / "p.dxf")
    b = dxf_io.read_dxf(tmp_path / "p.dxf")
    assert all(math.isnan(p.z) for p in b.points)


# ============================================================ GIS
@pytest.mark.parametrize("ext,driver", [(".gpkg", "GPKG"), (".geojson", "GeoJSON"), (".shp", "ESRI Shapefile")])
def test_gis_roundtrip(tmp_path, ext, driver):
    pr = make_site(4)
    pr.add_polyline([[0, 0, 5], [10, 0, 5], [10, 10, 5]], "ROAD")
    pr.add_polyline([[0, 0, 0], [10, 0, 0], [10, 10, 0]], "BLDG", closed=True, kind="parcel")
    f = tmp_path / ("out" + ext)
    st = gis_io.write_gis(f, pr, driver=driver, crs="EPSG:2276")
    assert st["points"] == 16 and st["lines"] == 1 and st["polygons"] == 1
    files = st["files"]
    assert (len(files) == 1) == (driver == "GPKG")
    pts_file = [x for x in files if "points" in x or driver == "GPKG"][0]
    layer = "points" if driver == "GPKG" else None
    b = gis_io.read_gis(pts_file, layer=layer)
    assert len(b.points) == 16 and "2276" in str(b.info["crs"])
    p0 = next(iter(pr.points.values()))
    q0 = b.points[0]
    assert (q0.x, q0.y, q0.z) == pytest.approx((p0.x, p0.y, p0.z), abs=1e-6)
    assert q0.number == p0.number and q0.desc == p0.desc
    if driver == "GPKG":
        names = [n for n, _ in gis_io.list_gis_layers(f)]
        assert {"points", "lines", "polygons"} <= set(names)
        bl = gis_io.read_gis(f, layer="polygons")
        assert len(bl.polylines) == 1 and bl.polylines[0].closed
        assert bl.polylines[0].attrs["area"] == pytest.approx(50.0)


def test_gis_reproject_on_export(tmp_path):
    pr = make_site(3)
    tf = C.make_transform(2276, 4326)
    f = tmp_path / "ll.geojson"
    gis_io.write_gis(f, pr, driver="GeoJSON", xy_fn=lambda x, y: tf(x, y), crs="EPSG:4326", include_lines=False)
    b = gis_io.read_gis(tmp_path / "ll_points.geojson")
    assert -97.5 < b.points[0].x < -96.0 and 32.0 < b.points[0].y < 33.5


def test_gis_import_polygons_and_field_guess(tmp_path):
    import shapely
    from pyogrio import raw
    geoms = shapely.to_wkb(np.array([shapely.Polygon([(0, 0), (10, 0), (10, 10), (0, 10)],
                                                     holes=[[(2, 2), (4, 2), (4, 4), (2, 4)]]),
                                     shapely.Point(5, 5, 7)], dtype=object), flavor="iso")
    f = tmp_path / "m.gpkg"
    raw.write(str(f), geometry=geoms, field_data=[np.array(["a", "b"], dtype=object), np.array([1.0, 2.0])],
              fields=["Name", "Elev"], crs="EPSG:2276", geometry_type="Unknown", driver="GPKG", layer="mix")
    b = gis_io.read_gis(f, layer="mix")
    assert gis_io.guess_fields(b.info["fields"]) == {"number": "Name", "elevation": "Elev"}
    assert len(b.polylines) == 2 and [p.attrs.get("hole") for p in b.polylines] == [None, True]
    assert b.points[0].number == "b" and b.points[0].z in (2.0, 7.0)


# ============================================================ KML
def test_kml_export_import_and_kmz(tmp_path):
    pr = make_site(3)
    pr.add_polyline([[2552000, 6967000, 0], [2552100, 6967100, 0]], "ROAD")
    tf = C.make_transform(2276, 4326)
    f = tmp_path / "o.kml"
    st = kml_io.write_kml(f, pr, lambda x, y: tf(x, y))
    assert st["points"] == 9 and st["lines"] == 1
    b = kml_io.read_kml(f)
    assert b.info["geographic"] and len(b.points) == 9 and len(b.polylines) == 1
    assert b.points[0].number == "1" and -96.9 < b.points[0].x < -96.5 and 32.5 < b.points[0].y < 33.0
    k = tmp_path / "o.kmz"
    kml_io.write_kml(k, pr, lambda x, y: tf(x, y))
    assert "doc.kml" in zipfile.ZipFile(k).namelist()
    assert len(kml_io.read_kml(k).points) == 9


def test_kml_import_groundoverlay_and_google_earth_style(tmp_path):
    kml = """<?xml version="1.0"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>
    <Folder><name>Checks</name><Placemark><name>101</name><description><![CDATA[<b>edge</b> of pavement]]></description>
      <Point><coordinates>-96.5992,32.7668,0</coordinates></Point></Placemark></Folder>
    <GroundOverlay><name>Plan</name><Icon><href>files/plan.png</href></Icon>
      <LatLonBox><north>32.78</north><south>32.77</south><east>-96.59</east><west>-96.60</west><rotation>5</rotation></LatLonBox></GroundOverlay>
    </Document></kml>"""
    z = tmp_path / "ge.kmz"
    import io as _io
    from PIL import Image
    buf = _io.BytesIO()
    Image.new("RGB", (4, 4), "red").save(buf, "PNG")
    with zipfile.ZipFile(z, "w") as zz:
        zz.writestr("doc.kml", kml)
        zz.writestr("files/plan.png", buf.getvalue())
    b = kml_io.read_kml(z, cache_dir=tmp_path)
    assert b.points[0].number == "101" and b.points[0].desc == "edge of pavement" and b.points[0].layer == "Checks"
    ov = b.info["overlays"][0]
    assert ov["north"] == 32.78 and ov["rotation"] == 5.0 and ov["file"] and open(ov["file"], "rb").read(4) == b"\x89PNG"


# ============================================================ reports
def test_reports_render_all_formats(tmp_path):
    pr = make_site(8)
    sf, _ = build_surface_from_project(pr, "EG", {})
    pr.add_surface(sf)
    from plumbline.core.qa import run_checks
    vol = sf.tin().volume_to_datum(sf.pts[:, 2].min())
    reps = [reports.points_report(pr), reports.surface_report(pr, sf), reports.volume_report(pr, vol),
            reports.polyline_report(pr, None), reports.qa_report(pr, run_checks(pr)),
            reports.crs_report(pr, "auto (2 m)")]
    pr.add_polyline([[0, 0, 0], [10, 0, 0], [10, 10, 0]], "BLDG", closed=True, bulges=[0, 0.5, 0])
    reps[3] = reports.polyline_report(pr, None)
    for i, r in enumerate(reps):
        h = reports.to_html(r)
        assert "<html>" in h and html.escape(r.title) in h
        reports.to_csv(r, tmp_path / f"r{i}.csv")
        reports.to_xlsx(r, tmp_path / f"r{i}.xlsx")
    # the Coordinate System Report carries the datum-shift line (what the removed imagery report
    # used to print): which operation PROJ picked, and its stated accuracy
    crs_html = reports.to_html(reps[5])
    assert "Datum transformation" in crs_html and "auto (2 m)" in crs_html
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "r4.xlsx")
    assert "Summary" in wb.sheetnames and "Data quality checks" in wb.sheetnames


def test_a_project_holding_legacy_imagery_checks_still_opens_and_saves(tmp_path):
    """Item 16 removed the checks, not the data: an old job must open, show and re-save unchanged."""
    from plumbline.core.model import ImageryCheck
    pr = make_site(4)
    for i, p in enumerate(list(pr.points.values()), start=1):
        pr.checks[i] = ImageryCheck(i, p.id, p.number, (p.x, p.y), (p.x + 1.0, p.y - 0.5), "Esri World Imagery",
                                    "an office note", "2026-01-02 09:30", "thumbdata", i % 2 == 0)
    path = tmp_path / "legacy.plb"
    pr.save(path)
    back = Project.load(path)
    assert len(back.checks) == len(pr.checks)
    c = back.checks[2]
    assert (c.layer_name, c.note, c.stamp, c.thumb) == ("Esri World Imagery", "an office note", "2026-01-02 09:30", "thumbdata")
    assert [x.include for x in back.checks.values()] == [x.include for x in pr.checks.values()]
    back.save()                                   # and saving it again keeps them


# ------------------------------------------------------------------ review fixes (v0.1 -> v0.2)
def test_write_dxf_takes_the_path_first():
    """Every writer in plumbline.io is (path, project, ...) - write_dxf used to be the odd one."""
    import inspect
    sig = list(inspect.signature(dxf_io.write_dxf).parameters)
    assert sig[:2] == ["path", "project"], sig


def test_write_dxf_rejects_the_old_argument_order(tmp_path):
    """A swapped call must fail loudly and helpfully, not deep inside with an AttributeError."""
    pr = Project("Swapped")
    pr.add_point(0, 0, 0, "1")
    with pytest.raises(TypeError) as e:
        dxf_io.write_dxf(pr, tmp_path / "x.dxf")
    assert "path comes first" in str(e.value) or "write_dxf(path, project" in str(e.value)
    # and the documented order still works
    st = dxf_io.write_dxf(tmp_path / "x.dxf", pr)
    assert st["points"] == 1 and (tmp_path / "x.dxf").exists()


def test_landxml_full_circle_curve(tmp_path):
    """A closed Curve (Start == End) is a full circle - it must not blow up the import."""
    f = tmp_path / "circle.xml"
    f.write_text("""<LandXML xmlns="http://www.landxml.org/schema/LandXML-1.2" version="1.2">
 <Units><Imperial linearUnit="USSurveyFoot"/></Units>
 <Parcels><Parcel name="Round"><CoordGeom>
   <Curve rot="ccw" crvType="arc" radius="10"><Start>10 0</Start><Center>0 0</Center><End>10 0</End></Curve>
 </CoordGeom></Parcel></Parcels>
</LandXML>""")
    b = landxml.read_landxml(f)
    assert len(b.polylines) == 1
    pl = b.polylines[0]
    assert pl.closed and len(pl.verts) == 2 and pl.bulges is not None
    assert np.allclose(pl.bulges, 1.0)
    # the shape really is a circle of radius 10 centred on the origin
    from plumbline.core import geometry as G
    assert G.polygon_area(pl.verts, pl.bulges) == pytest.approx(math.pi * 100.0, rel=1e-9)
    assert G.polyline_length(pl.verts, pl.bulges, True) == pytest.approx(2 * math.pi * 10.0, rel=1e-9)


def test_csv_header_coordinate_pair_is_trusted_as_is(tmp_path):
    """A header naming E and N must not have extra roles invented for it."""
    f = tmp_path / "h.csv"
    f.write_text("East,North,Elev,Desc\n"
                 "5000.0,3000.0,100.0,GS\n"
                 "5010.0,3010.0,101.0,EP\n")
    sn = CSV.sniff(f)
    assert sn.has_header
    assert "easting" in sn.roles and "northing" in sn.roles
    m = CSV.CsvMapping(delimiter=sn.delimiter, skip_rows=1, roles=sn.roles)
    b = CSV.read_points(f, m)
    assert len(b.points) == 2
    assert (b.points[0].x, b.points[0].y) == (5000.0, 3000.0)      # easting first, as mapped
    assert b.points[0].desc == "GS"
